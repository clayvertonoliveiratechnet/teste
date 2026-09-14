import json, math, re, unicodedata, uuid
from datetime import date, datetime, timezone
import openpyxl

ARQ = r'C:\Users\User\Desktop\TechNET-Rotas-MDU\fonte_live.xlsx'
SAIDA = r'C:\Users\User\Desktop\TechNET-Rotas-MDU\app\migracao_rotas.sql'
NS = uuid.UUID('2ad97764-e048-4b32-a0dd-27166003064f')
wb = openpyxl.load_workbook(ARQ, read_only=True, data_only=True)

def sem_acento(v):
    return ''.join(c for c in unicodedata.normalize('NFD', str(v or '')) if unicodedata.category(c) != 'Mn')

def cidade(v):
    s = sem_acento(v).upper()
    for chave, ident in [('NATAL','NATAL'),('MOSSORO','MOSSORO'),('FORTALEZA','FORTALEZA'),('RECIFE','RECIFE')]:
        if chave in s: return ident
    return None

def texto(v):
    if v is None: return None
    s = str(v).strip()
    return s or None

def inteiro(v):
    try: return max(0, int(round(float(v or 0))))
    except: return 0
def codigo(v):
    if v is None: return None
    if isinstance(v, int): return str(v)
    if isinstance(v, float):
        if math.isfinite(v) and v.is_integer(): return str(int(v))
        return format(v, '.15g')
    s = str(v).strip()
    return s[:-2] if re.fullmatch(r'\d+\.0', s) else (s or None)

def data_iso(v):
    if isinstance(v, datetime): return v.date().isoformat()
    if isinstance(v, date): return v.isoformat()
    s = texto(v)
    if not s: return None
    for fmt in ('%d/%m/%Y','%d/%m','%Y-%m-%d'):
        try:
            d = datetime.strptime(s, fmt)
            if fmt == '%d/%m': d = d.replace(year=2026)
            return d.date().isoformat()
        except: pass
    return None

def instante(v):
    if isinstance(v, datetime):
        return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat().replace('+00:00','Z')
    if isinstance(v, date):
        return datetime(v.year,v.month,v.day,tzinfo=timezone.utc).isoformat().replace('+00:00','Z')
    return texto(v)

def seguro(v):
    return v.isoformat() if isinstance(v, (datetime,date)) else v
geral = []
ws = wb['Geral']
for linha, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    cid = cidade(r[0])
    if not cid or not any(v not in (None,'') for v in r[:16]): continue
    legado = {
        'idCardTrello':seguro(r[16]), 'linkCardTrello':seguro(r[17]),
        'listaCardTrello':seguro(r[18]), 'enviadoTrelloEm':seguro(r[19]),
        'liberadoEm':seguro(r[20]), 'atualizadoTrelloEm':seguro(r[21]),
        'statusEnviadoTrello':seguro(r[22]), 'prazoCardTrello':seguro(r[23]),
        'hashCardTrello':seguro(r[24]), 'sourceRow':linha
    }
    geral.append({
        'id':str(uuid.uuid5(NS, f'geral:{linha}')), 'city_id':cid,
        'team':texto(r[1]), 'support':bool(r[2]) if isinstance(r[2],bool) else bool(texto(r[2])),
        'support_name':texto(r[3]), 'project_type':texto(r[4]), 'synergy':texto(r[5]),
        'property_code':codigo(r[6]), 'address':texto(r[7]) or '', 'name':texto(r[8]),
        'hps':inteiro(r[9]), 'construction_date':data_iso(r[10]), 'month_label':texto(r[11]),
        'status':texto(r[12]), 'inspection':texto(r[13]), 'measured':texto(r[14]),
        'notes':texto(r[15]), 'legacy_data':legado, 'source':'PLANILHA'
    })

rota = []
ws = wb['ADMINISTRATIVO']
atual = None
rota_data = '2026-09-14'
for linha, r in enumerate(ws.iter_rows(values_only=True), start=1):
    a = texto(r[0]) or ''
    b = texto(r[1]) or ''
    if 'ROTA DIÁRIA' in a.upper() or 'ROTA DIARIA' in sem_acento(a).upper():
        m = re.search(r'(\d{2}/\d{2}/\d{4})', a)
        if m: rota_data = data_iso(m.group(1)) or rota_data
        continue
    nova = cidade(a)
    if nova and (' - ' in a or '|' in a):
        atual = nova
        continue
    if not atual: continue
    if sem_acento(b).upper() == 'TECNICO' or sem_acento(a).upper() == 'MAQUINA DE FUSAO': continue
    if sem_acento(a).upper() == 'TOTAL' or sem_acento(b).upper() == 'TOTAL':
        atual = None
        continue
    tecnico, auxiliar, endereco = b, texto(r[2]) or '', texto(r[3]) or ''
    if not tecnico and not auxiliar: continue
    status = texto(r[6]) or 'EM CONSTRUÇÃO'
    m = re.match(r'^(\d{6,12})', endereco)
    chave = f'{rota_data}|{atual}|{tecnico}|{auxiliar}|{endereco}'
    rota.append({
        'id':str(uuid.uuid5(NS, 'rota:'+chave)), 'city_id':atual,
        'fusion_machine':bool(r[0]) if isinstance(r[0], bool) else False,
        'technician':tecnico, 'assistant':auxiliar, 'property_code':m.group(1) if m else None,
        'address':endereco, 'hps':inteiro(r[4]), 'blocks':inteiro(r[5]), 'status':status,
        'start_date':data_iso(r[7]), 'synergy':texto(r[8]), 'project_type':texto(r[9]),
        'notes':None, 'route_date':rota_data
    })

historico = []
ws = wb['HISTORICO_ROTA']
for linha, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
    cid, status = cidade(r[1]), texto(r[8])
    if not cid or not status: continue
    historico.append({
        'id':str(uuid.uuid5(NS, f'historico:{linha}')),
        'occurred_at':instante(r[0]) or datetime.now(timezone.utc).isoformat(),
        'city_id':cid, 'route_date':data_iso(r[2]), 'source_row':inteiro(r[3]) or None,
        'team':texto(r[4]), 'address':texto(r[5]), 'hps':inteiro(r[6]), 'blocks':inteiro(r[7]),
        'status':status, 'start_date':data_iso(r[9]), 'synergy':texto(r[10]), 'notes':texto(r[11])
    })

fila = []
ws = wb['PROXIMAS ROTAS']
for linha, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    cid, endereco = cidade(r[0]), texto(r[1])
    if not cid or not endereco: continue
    try: prioridade = max(1, min(3, int(float(r[3] or 2))))
    except: prioridade = 2
    ativo = sem_acento(r[8]).upper() not in ('FALSE','NAO','0')
    fila.append({
        'id':str(uuid.uuid5(NS, f'fila:{linha}')), 'city_id':cid, 'address':endereco,
        'status':texto(r[2]) or 'PENDENTE', 'priority':prioridade, 'notes':texto(r[4]),
        'latitude':r[5] if isinstance(r[5],(int,float)) else None,
        'longitude':r[6] if isinstance(r[6],(int,float)) else None,
        'assigned_to':texto(r[7]), 'active':ativo
    })
def bloco_insert(tabela, linhas, tipos, lote=250):
    if not linhas: return ''
    colunas = [c for c,_ in tipos]
    cols = ', '.join(colunas)
    defs = ', '.join(f'{c} {t}' for c,t in tipos)
    partes = []
    for i in range(0, len(linhas), lote):
        payload = json.dumps([{c:r.get(c) for c in colunas} for r in linhas[i:i+lote]], ensure_ascii=False, separators=(',',':'))
        partes.append(
            f"insert into public.{tabela} ({cols})\n"
            f"select {cols} from jsonb_to_recordset($data${payload}$data$::jsonb) as x({defs});\n"
        )
    return ''.join(partes)

tipos_geral = [('id','uuid'),('city_id','text'),('team','text'),('support','boolean'),('support_name','text'),('project_type','text'),('synergy','text'),('property_code','text'),('address','text'),('name','text'),('hps','integer'),('construction_date','date'),('month_label','text'),('status','text'),('inspection','text'),('measured','text'),('notes','text'),('legacy_data','jsonb'),('source','text')]
tipos_rota = [('id','uuid'),('city_id','text'),('fusion_machine','boolean'),('technician','text'),('assistant','text'),('property_code','text'),('address','text'),('hps','integer'),('blocks','integer'),('status','text'),('start_date','date'),('synergy','text'),('project_type','text'),('notes','text'),('route_date','date')]
tipos_hist = [('id','uuid'),('occurred_at','timestamptz'),('city_id','text'),('route_date','date'),('source_row','integer'),('team','text'),('address','text'),('hps','integer'),('blocks','integer'),('status','text'),('start_date','date'),('synergy','text'),('notes','text')]
tipos_fila = [('id','uuid'),('city_id','text'),('address','text'),('status','text'),('priority','smallint'),('notes','text'),('latitude','double precision'),('longitude','double precision'),('assigned_to','text'),('active','boolean')]
sql = ['begin;\n']
sql.append("delete from public.route_general where source='PLANILHA';\n")
sql.append('delete from public.route_daily;\n')
sql.append('delete from public.route_history;\n')
sql.append('delete from public.route_queue;\n')
sql.append(bloco_insert('route_general', geral, tipos_geral))
sql.append(bloco_insert('route_daily', rota, tipos_rota))
sql.append(bloco_insert('route_history', historico, tipos_hist))
sql.append(bloco_insert('route_queue', fila, tipos_fila))
sql.append('commit;\n')
with open(SAIDA, 'w', encoding='utf-8') as f:
    f.write(''.join(sql))

print(json.dumps({
    'geral':len(geral), 'rota':len(rota), 'historico':len(historico), 'fila':len(fila),
    'rota_data':rota_data,
    'hps_natal':sum(r['hps'] for r in rota if r['city_id']=='NATAL'),
    'hps_recife':sum(r['hps'] for r in rota if r['city_id']=='RECIFE'),
    'arquivo_sql':SAIDA
}, ensure_ascii=False))
